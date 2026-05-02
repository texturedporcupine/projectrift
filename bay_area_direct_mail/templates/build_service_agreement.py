"""
Build the California Advertising Services Agreement as a .docx.

Reads `templates/service_agreement_ca.md` (markdown source of truth) and
emits a clean Word document with bolded headings, normal body, and merge
placeholders preserved (`{{variable}}`) so the operator can fill them in.

Usage:
    python templates/build_service_agreement.py
        # writes templates/service_agreement_ca.docx

    python templates/build_service_agreement.py --advertiser "Sycamore Valley Dental" \
        --community danville --total-fee 700
        # writes templates/agreement_sycamore-valley-dental_danville.docx
        # (with the named variables pre-filled)
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


HERE = Path(__file__).parent
SRC_MD = HERE / "service_agreement_ca.md"


def slugify(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return s or "advertiser"


def add_heading(doc, text, *, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 0:
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0x0B, 0x1F, 0x3A)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x0B, 0x1F, 0x3A)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x32, 0x44, 0x66)
    return p


def add_para(doc, text, *, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        run = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run.text = part[1:-1]
            run.italic = True
        else:
            run.text = part
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0x32, 0x44, 0x66)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
    return p


def render_md_to_docx(md_text: str, doc: Document, replacements: dict[str, str]):
    """Lightweight markdown -> docx renderer for the agreement template."""
    for k, v in replacements.items():
        md_text = md_text.replace("{{" + k + "}}", v)

    lines = md_text.splitlines()
    i = 0
    in_list = False
    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip()

        if not stripped:
            in_list = False
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("# "):
            add_heading(doc, stripped[2:].strip(), level=0)
        elif stripped.startswith("## "):
            add_heading(doc, stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), level=2)
        elif stripped.startswith("> "):
            add_para(doc, stripped[2:].strip(), italic=True, size=10)
        elif stripped == "---":
            doc.add_paragraph("─" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("- "):
            in_list = True
            p = doc.add_paragraph(style="List Bullet")
            for part in re.split(r"(\*\*[^*]+\*\*)", stripped[2:].strip()):
                if not part:
                    continue
                run = p.add_run()
                if part.startswith("**") and part.endswith("**"):
                    run.text = part[2:-2]
                    run.bold = True
                else:
                    run.text = part
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x32, 0x44, 0x66)
        else:
            add_para(doc, stripped)
        i += 1


def build_agreement(out_path: Path, replacements: dict[str, str]) -> Path:
    md_text = SRC_MD.read_text(encoding="utf-8")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    render_md_to_docx(md_text, doc, replacements)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="Build CA Advertising Services Agreement (.docx)")
    ap.add_argument("--advertiser", default=None,
                    help="Advertiser legal name (used in filename + body).")
    ap.add_argument("--community", default=None,
                    help="Community name (e.g., 'Danville')")
    ap.add_argument("--zip", default=None,
                    help="Campaign zip code(s)")
    ap.add_argument("--total-fee", type=int, default=None,
                    help="Total slot fee in USD (e.g., 700)")
    ap.add_argument("--out", default=None,
                    help="Override output path")
    args = ap.parse_args()

    replacements: dict[str, str] = {
        # Defaults left as merge tokens so the operator (or a Cowork run)
        # fills them in. Override with CLI args for partially-filled output.
        "effective_date": "{{effective_date}}",
        "provider_legal_name": "{{provider_legal_name}}",
        "provider_dba": "{{provider_dba}}",
        "provider_county": "{{provider_county}}",
        "provider_address": "{{provider_address}}",
        "advertiser_legal_name": "{{advertiser_legal_name}}",
        "advertiser_address": "{{advertiser_address}}",
        "community": "{{community}}",
        "zip_codes": "{{zip_codes}}",
        "drop_date": "{{drop_date}}",
        "advertiser_category": "{{advertiser_category}}",
        "total_fee": "{{total_fee}}",
        "deposit_amount": "{{deposit_amount}}",
        "balance_amount": "{{balance_amount}}",
        "provider_signer_name": "{{provider_signer_name}}",
        "provider_signer_title": "{{provider_signer_title}}",
        "provider_signer_date": "{{provider_signer_date}}",
        "advertiser_signer_name": "{{advertiser_signer_name}}",
        "advertiser_signer_title": "{{advertiser_signer_title}}",
        "advertiser_signer_date": "{{advertiser_signer_date}}",
    }

    if args.advertiser:
        replacements["advertiser_legal_name"] = args.advertiser
    if args.community:
        replacements["community"] = args.community
    if args.zip:
        replacements["zip_codes"] = args.zip
    if args.total_fee is not None:
        half = args.total_fee // 2
        replacements["total_fee"] = f"{args.total_fee:,}"
        replacements["deposit_amount"] = f"{half:,}"
        replacements["balance_amount"] = f"{args.total_fee - half:,}"

    if args.out:
        out = Path(args.out)
    elif args.advertiser and args.community:
        slug = f"{slugify(args.advertiser)}_{slugify(args.community)}"
        out = HERE / f"agreement_{slug}.docx"
    else:
        out = HERE / "service_agreement_ca.docx"

    saved = build_agreement(out, replacements)
    print(f"Wrote {saved}")


if __name__ == "__main__":
    main()
