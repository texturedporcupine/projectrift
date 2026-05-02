"""
Build the SF Bay Area prospect tracker xlsx.

Cloned from the Austin Avery Ranch template (`prospecttracker_averyranch78717.xlsx`).
Same column structure, same rating system, plus per-community sample
prospects so you can start from a primed sheet rather than a blank one.

Conventions:
    - One sheet per community.
    - Read all sheets at once with `pandas.read_excel(path, sheet_name=None)`.
    - Rating: IDEAL > STRONG > GOOD > MAYBE > SKIP.

Usage:
    python prospects/build_prospect_tracker.py
        # writes prospect-tracker_{community}-{zip}.xlsx for each community

    python prospects/build_prospect_tracker.py --community danville
        # only write Danville
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


# ---------------------------------------------------------------------------
# Tracker schema
# ---------------------------------------------------------------------------

COLUMNS = [
    ("Category",            18),
    ("Business Name",       28),
    ("Owner / Contact",     22),
    ("Phone",               16),
    ("Email",               28),
    ("Website",             26),
    ("Address",             32),
    ("Rating",              10),
    ("Category Locked?",    16),
    ("Outreach Stage",      18),
    ("Last Touch",          14),
    ("Next Action",         24),
    ("Quoted Price",        14),
    ("Deposit Received",    16),
    ("Notes",               40),
]

RATINGS = ["IDEAL", "STRONG", "GOOD", "MAYBE", "SKIP"]
LOCKED  = ["YES", "NO"]
STAGES  = [
    "Researched",
    "Email 1 sent",
    "Email 2 sent",
    "Email 3 sent",
    "Phone scheduled",
    "Phone done",
    "Meeting scheduled",
    "Proposal sent",
    "Verbal yes",
    "Signed (deposit pending)",
    "Booked (deposit received)",
    "Lost",
]

NAVY  = "FF0B1F3A"
GOLD  = "FFD4A84B"
WHITE = "FFFFFFFF"
LIGHT = "FFF4EFE6"
SLATE = "FF324466"
GREEN = "FFC7E5C5"
RED   = "FFE5C5C5"


@dataclass
class Prospect:
    category: str
    name: str
    contact: str = ""
    phone: str = ""
    email: str = ""
    website: str = ""
    address: str = ""
    rating: str = "STRONG"
    locked: str = "NO"
    stage: str = "Researched"
    notes: str = ""


@dataclass
class CommunitySheet:
    key: str
    label: str
    zip_code: str
    target_categories: list[str]
    seed_prospects: list[Prospect] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Seed prospects per community
#
# These are *category exemplars*, not a guarantee that the named businesses are
# best-fit prospects. The point is to start the operator with realistic
# search hooks (corridor + brand-style) rather than a blank sheet. Always
# verify in Google Maps before reaching out.
# ---------------------------------------------------------------------------

DANVILLE = CommunitySheet(
    key="danville",
    label="Danville",
    zip_code="94526",
    target_categories=[
        "HVAC", "Pest Control", "General Dentistry", "Pediatric Dentistry",
        "Pool Service", "Landscape Design", "Plumbing", "Roofing",
        "House Cleaning", "Med Spa / Aesthetics",
    ],
    seed_prospects=[
        Prospect("HVAC", "Tri-Valley Heating & Air (example)",
                 contact="Owner",
                 address="Crow Canyon Rd corridor, Danville",
                 rating="IDEAL",
                 notes="High-ticket equipment swaps; HOA neighborhoods skew older HVAC."),
        Prospect("Pest Control", "Diablo Pest Solutions (example)",
                 address="Hartz Ave corridor, Danville",
                 rating="STRONG",
                 notes="Quarterly contracts = recurring rev; Blackhawk skews well."),
        Prospect("General Dentistry", "Sycamore Valley Dental (example)",
                 address="Camino Tassajara, Danville",
                 rating="IDEAL",
                 notes="Family practice; mention Blackhawk school families."),
        Prospect("Pediatric Dentistry", "Diablo Pediatric Dental (example)",
                 address="San Ramon Valley Blvd, Danville",
                 rating="IDEAL",
                 notes="Excellent EDDM ROI; targets young families w/ kids."),
        Prospect("Pool Service", "Blackhawk Pool & Spa (example)",
                 address="Blackhawk corridor, Danville",
                 rating="STRONG",
                 notes="Many backyard pools in 94526; weekly service plans."),
        Prospect("Landscape Design", "Diablo Landscape Architects (example)",
                 address="Diablo Rd, Danville",
                 rating="IDEAL",
                 notes="Premium landscape design — Tier 1 placement."),
        Prospect("Plumbing", "Sycamore Plumbing Co (example)",
                 address="San Ramon Valley Blvd",
                 rating="GOOD",
                 notes="Established homes = repipes / water heater swaps."),
        Prospect("Roofing", "Tri-Valley Roof Co (example)",
                 address="Camino Ramon, Danville",
                 rating="GOOD",
                 notes="Composition + tile re-roofs; 25-yr cycle hits often here."),
        Prospect("House Cleaning", "Blackhawk Cleaning Concierge (example)",
                 address="Blackhawk Plaza, Danville",
                 rating="STRONG",
                 notes="Premium recurring service; high % of dual-income HHs."),
        Prospect("Med Spa / Aesthetics", "Diablo Aesthetics (example)",
                 address="Hartz Ave, Danville",
                 rating="IDEAL",
                 notes="Best Tier 1 vertical here — premium placement target."),
    ],
)

SAN_RAMON = CommunitySheet(
    key="san_ramon",
    label="San Ramon",
    zip_code="94582-94583",
    target_categories=[
        "HVAC", "Pest Control", "Family Dentistry", "Pediatric Dentistry",
        "Tutoring / Test Prep", "Pool Service", "Landscape", "Plumbing",
        "House Cleaning", "Orthodontics",
    ],
    seed_prospects=[
        Prospect("HVAC", "Dougherty Valley Heating (example)",
                 address="Crow Canyon Rd, San Ramon",
                 rating="IDEAL",
                 notes="Master-planned tracts hitting 15-yr HVAC cycle."),
        Prospect("Pest Control", "Bishop Ranch Pest (example)",
                 address="Bishop Ranch corridor, San Ramon",
                 rating="STRONG"),
        Prospect("Family Dentistry", "Windemere Family Dental (example)",
                 address="Bollinger Canyon Rd, San Ramon",
                 rating="IDEAL",
                 notes="Family-heavy zip; SRVUSD parents."),
        Prospect("Pediatric Dentistry", "San Ramon Pediatric Dental (example)",
                 address="Alcosta Blvd, San Ramon",
                 rating="IDEAL"),
        Prospect("Tutoring / Test Prep", "Dougherty Valley Tutoring (example)",
                 address="Bollinger Canyon Rd, San Ramon",
                 rating="IDEAL",
                 notes="STEM-heavy parent demographic; very strong ROI."),
        Prospect("Pool Service", "San Ramon Pool Pros (example)",
                 address="Crow Canyon Rd, San Ramon",
                 rating="STRONG"),
        Prospect("Landscape", "Tri-Valley Landscape Co (example)",
                 address="Camino Ramon, San Ramon",
                 rating="STRONG"),
        Prospect("Plumbing", "Bishop Ranch Plumbing (example)",
                 address="Bishop Ranch corridor",
                 rating="GOOD"),
        Prospect("House Cleaning", "Windemere Home Care (example)",
                 address="Windemere, San Ramon",
                 rating="STRONG"),
        Prospect("Orthodontics", "San Ramon Orthodontics (example)",
                 address="Crow Canyon Rd",
                 rating="STRONG",
                 notes="Strong vertical; SRVUSD parents = high braces demand."),
    ],
)

PLEASANTON = CommunitySheet(
    key="pleasanton",
    label="Pleasanton",
    zip_code="94566-94588",
    target_categories=[
        "Pool Service", "Landscape", "Orthodontics", "General Dentistry",
        "House Cleaning", "HVAC", "Pest Control", "Med Spa", "Plumbing",
        "Roofing",
    ],
    seed_prospects=[
        Prospect("Pool Service", "Ruby Hill Pool Co (example)",
                 address="Ruby Hill, Pleasanton",
                 rating="IDEAL"),
        Prospect("Landscape", "Pleasanton Landscape Design (example)",
                 address="Bernal Ave",
                 rating="STRONG"),
        Prospect("Orthodontics", "Tri-Valley Orthodontics (example)",
                 address="Hopyard Rd",
                 rating="STRONG"),
        Prospect("General Dentistry", "Main Street Dental Pleasanton (example)",
                 address="Main St downtown",
                 rating="IDEAL"),
        Prospect("House Cleaning", "Castlewood Cleaning (example)",
                 address="Castlewood, Pleasanton",
                 rating="STRONG"),
        Prospect("HVAC", "Pleasanton Heating & Air (example)",
                 address="Santa Rita Rd",
                 rating="STRONG"),
        Prospect("Pest Control", "Pleasanton Pest Solutions (example)",
                 address="Hopyard Rd",
                 rating="GOOD"),
        Prospect("Med Spa", "Pleasanton Aesthetics (example)",
                 address="Main St downtown",
                 rating="IDEAL"),
        Prospect("Plumbing", "Tri-Valley Plumbing (example)",
                 address="Bernal Ave",
                 rating="GOOD"),
        Prospect("Roofing", "Pleasanton Roofing Co (example)",
                 address="Santa Rita Rd",
                 rating="GOOD"),
    ],
)

DUBLIN = CommunitySheet(
    key="dublin",
    label="Dublin",
    zip_code="94568",
    target_categories=[
        "HVAC", "Pest Control", "Tutoring / Test Prep", "Family Dentistry",
        "Pediatric Dentistry", "House Cleaning", "Pool Service",
        "Landscape", "Plumbing", "Med Spa",
    ],
    seed_prospects=[
        Prospect("HVAC", "Dublin Heating & Air (example)",
                 address="Dublin Blvd",
                 rating="STRONG"),
        Prospect("Tutoring / Test Prep", "Dublin Ranch Tutoring (example)",
                 address="Tassajara Rd",
                 rating="IDEAL"),
        Prospect("Pediatric Dentistry", "Dublin Kids Dental (example)",
                 address="Hacienda Dr",
                 rating="IDEAL"),
    ],
)

LOS_GATOS = CommunitySheet(
    key="los_gatos",
    label="Los Gatos",
    zip_code="95030-95032",
    target_categories=[
        "Med Spa / Aesthetics", "Luxury Home Renovation", "Pool Service",
        "General / Cosmetic Dentistry", "Premium Landscape",
        "Concierge Home Services", "Pest Control", "HVAC", "Plumbing",
        "Roofing",
    ],
    seed_prospects=[
        Prospect("Med Spa / Aesthetics", "Los Gatos Aesthetics (example)",
                 address="Santa Cruz Ave",
                 rating="IDEAL",
                 notes="Top Tier 1 vertical; premium placement target."),
        Prospect("Luxury Home Renovation", "Los Gatos Custom Homes (example)",
                 address="Los Gatos Blvd",
                 rating="IDEAL"),
        Prospect("Premium Landscape", "Los Gatos Landscape Architects (example)",
                 address="University Ave",
                 rating="STRONG"),
        Prospect("Cosmetic Dentistry", "Los Gatos Smile Studio (example)",
                 address="Santa Cruz Ave",
                 rating="IDEAL"),
    ],
)

SARATOGA = CommunitySheet(
    key="saratoga",
    label="Saratoga",
    zip_code="95070",
    target_categories=[
        "Luxury Home Renovation", "Med Spa", "Estate Planning",
        "Premium Landscape Architect", "Fine Jewelry",
        "Wine Retail", "Concierge Medicine", "Cosmetic Dentistry",
        "Custom Pool", "Auto Detailing (luxury)",
    ],
    seed_prospects=[
        Prospect("Luxury Home Renovation", "Saratoga Custom Homes (example)",
                 address="Big Basin Way",
                 rating="IDEAL"),
        Prospect("Med Spa", "Saratoga Aesthetics (example)",
                 address="Big Basin Way",
                 rating="IDEAL"),
        Prospect("Estate Planning", "Saratoga Estate Law (example)",
                 address="Saratoga Ave",
                 rating="STRONG"),
    ],
)

ALL_COMMUNITIES = [DANVILLE, SAN_RAMON, PLEASANTON, DUBLIN, LOS_GATOS, SARATOGA]


# ---------------------------------------------------------------------------
# Workbook construction
# ---------------------------------------------------------------------------

THIN = Side(style="thin", color="FFD8CEB6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(cell):
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center",
                               wrap_text=True)
    cell.border = BORDER


def style_body(cell):
    cell.font = Font(size=10, color=SLATE)
    cell.alignment = Alignment(horizontal="left", vertical="top",
                               wrap_text=True)
    cell.border = BORDER


def add_summary_sheet(wb: Workbook, communities: list[CommunitySheet]):
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90

    title = ws.cell(row=1, column=1,
                    value="SF Bay Area — Prospect Tracker")
    title.font = Font(bold=True, size=18, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

    rows = [
        ("Adapted from", "Avery Ranch (78717) Austin TX prospect tracker template."),
        ("Conventions",  "One sheet per community. Read all sheets with "
                         "pandas.read_excel(path, sheet_name=None)."),
        ("Ratings",      "IDEAL > STRONG > GOOD > MAYBE > SKIP."),
        ("Stages",       "Researched -> Email 1 -> Email 2 -> Email 3 -> "
                         "Phone scheduled -> Phone done -> Meeting scheduled "
                         "-> Proposal sent -> Verbal yes -> Signed -> Booked / Lost."),
        ("Category lock", "When you sign a category, set Category Locked? = YES "
                          "and stop pursuing other businesses in that category."),
        ("Sheets in this workbook",
         ", ".join(f"{c.label} ({c.zip_code})" for c in communities)),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        a = ws.cell(row=i, column=1, value=k)
        a.font = Font(bold=True, color=SLATE, size=11)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=i, column=2, value=v)
        b.font = Font(color=SLATE, size=11)
        b.alignment = Alignment(wrap_text=True, vertical="top")

    ws.cell(row=len(rows) + 5, column=1,
            value="To add a new community sheet:").font = Font(
        bold=True, color=NAVY, size=12)
    ws.cell(row=len(rows) + 6, column=1,
            value="1. Duplicate any community sheet.").font = Font(color=SLATE)
    ws.cell(row=len(rows) + 7, column=1,
            value="2. Rename to the new community.").font = Font(color=SLATE)
    ws.cell(row=len(rows) + 8, column=1,
            value="3. Update the header rows with the new zip and target categories.").font = Font(color=SLATE)


def write_community_sheet(wb: Workbook, c: CommunitySheet):
    ws = wb.create_sheet(title=f"{c.label} {c.zip_code}"[:31])

    ws.cell(row=1, column=1,
            value=f"{c.label}, CA  ({c.zip_code})").font = Font(
        bold=True, size=16, color=WHITE)
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))

    sub = ws.cell(row=2, column=1,
                  value="Target categories: " + ", ".join(c.target_categories))
    sub.font = Font(size=10, italic=True, color=SLATE)
    sub.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
    ws.row_dimensions[2].height = 30

    header_row = 4
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        style_header(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 28

    first_data_row = header_row + 1
    last_data_row = first_data_row + max(len(c.seed_prospects), 30) - 1

    rating_dv = DataValidation(type="list",
                               formula1=f'"{",".join(RATINGS)}"',
                               allow_blank=True)
    locked_dv = DataValidation(type="list",
                               formula1=f'"{",".join(LOCKED)}"',
                               allow_blank=True)
    stage_dv = DataValidation(type="list",
                              formula1=f'"{",".join(STAGES)}"',
                              allow_blank=True)
    deposit_dv = DataValidation(type="list",
                                formula1=f'"{",".join(LOCKED)}"',
                                allow_blank=True)
    for dv in (rating_dv, locked_dv, stage_dv, deposit_dv):
        ws.add_data_validation(dv)
    rating_dv.add(f"H{first_data_row}:H{last_data_row}")
    locked_dv.add(f"I{first_data_row}:I{last_data_row}")
    stage_dv.add(f"J{first_data_row}:J{last_data_row}")
    deposit_dv.add(f"N{first_data_row}:N{last_data_row}")

    for i in range(first_data_row, last_data_row + 1):
        for col_idx in range(1, len(COLUMNS) + 1):
            style_body(ws.cell(row=i, column=col_idx))
        ws.row_dimensions[i].height = 32

    for i, p in enumerate(c.seed_prospects):
        r = first_data_row + i
        values = [p.category, p.name, p.contact, p.phone, p.email,
                  p.website, p.address, p.rating, p.locked, p.stage,
                  "", "", "", "", p.notes]
        for col_idx, v in enumerate(values, start=1):
            ws.cell(row=r, column=col_idx, value=v)

    ws.freeze_panes = ws.cell(row=first_data_row, column=1)


def build_workbook(communities: list[CommunitySheet], out_path: Path) -> Path:
    wb = Workbook()
    add_summary_sheet(wb, communities)
    for c in communities:
        write_community_sheet(wb, c)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def build_single_community_workbook(c: CommunitySheet, out_path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    write_community_sheet(wb, c)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(
        description="Build SF Bay Area prospect tracker xlsx files.")
    ap.add_argument("--community", default=None,
                    help=f"Build only this community ({', '.join(c.key for c in ALL_COMMUNITIES)})")
    ap.add_argument("--combined", action="store_true",
                    help="Also write a combined multi-community workbook.")
    args = ap.parse_args()
    base = Path(__file__).parent

    targets = ALL_COMMUNITIES
    if args.community:
        targets = [c for c in ALL_COMMUNITIES if c.key == args.community]
        if not targets:
            raise SystemExit(
                f"Unknown community '{args.community}'. "
                f"Choose from: {', '.join(c.key for c in ALL_COMMUNITIES)}")

    for c in targets:
        out = base / f"prospect-tracker_{c.key}-{c.zip_code}.xlsx"
        build_single_community_workbook(c, out)
        print(f"Wrote {out}")

    if args.combined or not args.community:
        out = base / "prospect-tracker_bay-area_all-communities.xlsx"
        build_workbook(ALL_COMMUNITIES, out)
        print(f"Wrote {out}")


if __name__ == "__main__":
    main()
