"""Generate the California Service Agreement template (.docx)."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path(__file__).resolve().parent.parent / "agreements" / "agreement_template_california.docx"


def H(doc, text, level=1):
    doc.add_heading(text, level=level)


def P(doc, text):
    p = doc.add_paragraph(text)
    return p


def main():
    doc = Document()
    doc.add_heading("Shared Postcard Direct Mail — Advertising Service Agreement", level=0)
    p = doc.add_paragraph()
    r = p.add_run("California / SF Bay Area · Template v1")
    r.italic = True

    H(doc, "1. Parties", 1)
    P(doc,
      "This Advertising Service Agreement (\"Agreement\") is entered into as of "
      "______________ (\"Effective Date\") between [YOUR BUSINESS NAME], a "
      "California sole proprietorship doing business as [DBA] (\"Publisher\"), "
      "and [ADVERTISER LEGAL NAME] (\"Advertiser\").")

    H(doc, "2. Services", 1)
    P(doc,
      "Publisher will produce a shared postcard advertising piece (\"Postcard\") "
      "and arrange for distribution to approximately [HOUSEHOLD COUNT] households "
      "in the [COMMUNITY NAME], CA target area (zip code(s) [ZIP]) via the United "
      "States Postal Service Every Door Direct Mail (EDDM) program.")
    P(doc,
      "Advertiser will receive one (1) exclusive ad slot in the business "
      "category of [CATEGORY] on the Postcard. No competing business in the same "
      "category will appear on the same Postcard or on any other Postcard "
      "distributed by Publisher to the same target area during the term of this "
      "Agreement.")

    H(doc, "3. Term", 1)
    P(doc,
      "This Agreement covers one (1) campaign with target in-home delivery on or "
      "about [TARGET DROP DATE]. Category exclusivity for Advertiser begins on "
      "the Effective Date and continues until the later of (a) thirty (30) days "
      "after target in-home delivery, or (b) the date Advertiser declines a "
      "renewal offer in writing.")

    H(doc, "4. Fees & Payment", 1)
    P(doc,
      "Total advertising fee: $[TOTAL]. Payment terms: 50% deposit due upon "
      "execution of this Agreement; remaining 50% due before the Postcard goes "
      "to print. All payments via ACH or check made payable to [YOUR BUSINESS "
      "NAME].")
    P(doc,
      "If Advertiser fails to pay the final 50% within five (5) business days of "
      "Publisher's request, Publisher may, at its option, replace Advertiser's "
      "ad slot, retain the deposit as liquidated damages, and pursue any other "
      "remedy available at law.")

    H(doc, "5. Creative & Approval", 1)
    P(doc,
      "Advertiser will provide logo files (vector preferred), brand colors, "
      "preferred offer/headline copy, photography (if applicable), and any "
      "regulatory disclosures within five (5) business days of the Effective "
      "Date. Publisher will produce a digital proof and Advertiser will have "
      "three (3) business days to request revisions or approve. Advertiser is "
      "responsible for the accuracy of all copy, claims, prices, and offers in "
      "its ad slot.")

    H(doc, "6. Distribution", 1)
    P(doc,
      "Publisher will arrange printing on minimum 14pt cardstock and submit the "
      "Postcards via USPS EDDM at the Destination Delivery Unit (DDU) serving "
      "the target zip code(s). Postage of $0.247/piece is included in the "
      "advertising fee. Actual delivery dates depend on USPS processing and may "
      "vary by 5–10 business days from the target drop date.")

    H(doc, "7. Performance", 1)
    P(doc,
      "Publisher does not guarantee any specific number of leads, calls, "
      "appointments, or sales. Direct-mail performance varies based on offer, "
      "creative, season, and Advertiser response capability. Publisher will, on "
      "request, provide aggregated campaign metrics (drop date, household count, "
      "route information).")

    H(doc, "8. Renewal & Right of First Refusal", 1)
    P(doc,
      "Advertiser will have a right of first refusal on the same category for "
      "the next campaign in the same target community at a discount of fifteen "
      "percent (15%) off the then-current rate, provided Advertiser exercises "
      "this right within fourteen (14) days after delivery of the current "
      "campaign.")

    H(doc, "9. Indemnification", 1)
    P(doc,
      "Advertiser agrees to indemnify and hold Publisher harmless from any "
      "claim arising out of the content of Advertiser's ad slot, including but "
      "not limited to claims of false advertising, intellectual property "
      "infringement, or regulatory violations.")

    H(doc, "10. Limitation of Liability", 1)
    P(doc,
      "Publisher's total liability under this Agreement is limited to the fees "
      "paid by Advertiser. In no event will Publisher be liable for indirect, "
      "incidental, consequential, or punitive damages.")

    H(doc, "11. Governing Law & Venue", 1)
    P(doc,
      "This Agreement is governed by the laws of the State of California, "
      "without regard to its conflict of laws rules. Venue for any dispute lies "
      "in the state or federal courts located in [COUNTY], California.")

    H(doc, "12. Entire Agreement", 1)
    P(doc,
      "This Agreement constitutes the entire understanding between the parties "
      "and supersedes all prior discussions. Amendments must be in writing and "
      "signed by both parties.")

    H(doc, "13. Signatures", 1)
    table = doc.add_table(rows=4, cols=2)
    table.style = "Light Grid"
    cells = table.rows[0].cells
    cells[0].text = "Publisher"
    cells[1].text = "Advertiser"
    cells = table.rows[1].cells
    cells[0].text = "Signature: _______________________"
    cells[1].text = "Signature: _______________________"
    cells = table.rows[2].cells
    cells[0].text = "Print name: ______________________"
    cells[1].text = "Print name: ______________________"
    cells = table.rows[3].cells
    cells[0].text = "Date: ____________________________"
    cells[1].text = "Date: ____________________________"

    p = doc.add_paragraph()
    r = p.add_run(
        "Publisher's California Fictitious Business Name (FBN) statement is on "
        "file with the [COUNTY] County Recorder, file no. [FBN-FILE-NO], filed "
        "[FBN-FILE-DATE].")
    r.italic = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
