"""Generate a market-neutral creative brief template (.docx)."""
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent.parent / "agreements" / "creative-brief_template.docx"


def H(doc, text, level=1):
    doc.add_heading(text, level=level)


def P(doc, text):
    doc.add_paragraph(text)


def kv_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Light Grid"
    for i, (k, v) in enumerate(rows):
        c = table.rows[i].cells
        c[0].text = k
        c[1].text = v
        c[0].width = c[0].width  # noqa: keep python-docx happy
    return table


def main():
    doc = Document()
    doc.add_heading("Postcard Creative Brief — Advertiser Slot", level=0)
    P(doc, "Fill out one brief per advertiser. We use this to design your ad slot.")

    H(doc, "1. Business Snapshot", 1)
    kv_table(doc, [
        ("Business name", ""),
        ("Owner / point of contact", ""),
        ("Phone", ""),
        ("Email", ""),
        ("Service area", ""),
        ("Website", ""),
        ("Years in business", ""),
        ("Google rating / # reviews", ""),
    ])

    H(doc, "2. Brand & Visual Assets", 1)
    P(doc, "Please attach (a) vector logo file (.ai/.eps/.svg), (b) brand color hex codes, "
           "(c) any photography you'd like featured, (d) staff photo if applicable.")
    kv_table(doc, [
        ("Primary brand color (hex)", ""),
        ("Secondary brand color (hex)", ""),
        ("Logo file format(s) provided", ""),
        ("Photo asset(s) provided", ""),
    ])

    H(doc, "3. The Offer", 1)
    P(doc, "What's the headline offer on this postcard? Be specific and time-bound when possible.")
    kv_table(doc, [
        ("Headline (max 8 words)", ""),
        ("Specific offer / discount", ""),
        ("Offer expiration date", ""),
        ("Promo code (optional)", ""),
        ("Required disclaimers / fine print", ""),
    ])

    H(doc, "4. Calls to Action & Tracking", 1)
    kv_table(doc, [
        ("Primary CTA (call / book / scan)", ""),
        ("Tracking phone number (we can issue one)", ""),
        ("Landing page URL (we can shortlink)", ""),
        ("QR code destination", ""),
    ])

    H(doc, "5. Differentiators", 1)
    P(doc, "Three reasons a homeowner should pick you over a competitor.")
    for i in range(1, 4):
        P(doc, f"{i}. ")

    H(doc, "6. What NOT to say", 1)
    P(doc, "Any words, claims, or imagery to avoid (e.g., regulatory restrictions, "
           "prior bad campaigns, competitor references).")

    H(doc, "7. Approval", 1)
    P(doc, "We will deliver a digital proof within 5 business days of receiving this "
           "brief plus all assets. You have 3 business days to approve or request revisions. "
           "Final balance is due before printing.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
