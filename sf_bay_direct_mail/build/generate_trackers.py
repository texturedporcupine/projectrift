"""Generate prospect tracker spreadsheets for Bay Area first-launch communities."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


OUT_DIR = Path(__file__).resolve().parent.parent / "trackers"

HEADER_FILL = PatternFill("solid", fgColor="0E1F3A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill("solid", fgColor="F4F1EA")
SUBHEADER_FONT = Font(bold=True, color="0E1F3A", size=11)
THIN = Side(border_style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CATEGORIES = [
    "HVAC",
    "Pest Control",
    "General Dentistry",
    "Pediatric Dentistry",
    "Pool Service",
    "Landscaping",
    "Plumbing",
    "Roofing",
    "House Cleaning",
    "Med Spa / Aesthetics",
]

RATINGS = ["IDEAL", "STRONG", "GOOD", "MAYBE", "PASS"]
STATUS = [
    "Not contacted",
    "Email 1 sent",
    "Email 2 sent",
    "Email 3 sent",
    "Phone follow-up",
    "Meeting booked",
    "Proposal sent",
    "Verbal yes",
    "Deposit paid",
    "Signed + paid in full",
    "Lost",
    "Renewal candidate",
]

PROSPECT_COLUMNS = [
    ("Category", 22),
    ("Business Name", 28),
    ("Website", 30),
    ("Phone", 16),
    ("Owner / Decision Maker", 24),
    ("Email", 28),
    ("Address", 32),
    ("Google Rating", 14),
    ("Years in Business", 14),
    ("Rating (IDEAL/STRONG/GOOD)", 22),
    ("Locked? (Y/N)", 12),
    ("Status", 22),
    ("Last Touch Date", 14),
    ("Next Action", 28),
    ("Quoted Price", 14),
    ("Deposit Paid?", 14),
    ("Notes", 50),
]

# Seed prospects per community. These are realistic templated examples to
# demonstrate the tracker layout — not customer data. Replace with real
# research before outreach.
SEED_PROSPECTS = {
    "Danville": [
        ("HVAC", "Diablo Valley Heating & Air", "diablovalleyhvac.example", "(925) 555-0140", "", "", "Danville, CA 94526", "", "", "IDEAL", "N", "Not contacted"),
        ("HVAC", "Tassajara Comfort Systems", "tassajaracomfort.example", "(925) 555-0142", "", "", "Danville, CA 94526", "", "", "STRONG", "N", "Not contacted"),
        ("Pest Control", "East Bay Pest Defense", "ebpestdefense.example", "(925) 555-0151", "", "", "Danville, CA 94526", "", "", "IDEAL", "N", "Not contacted"),
        ("Pool Service", "Sycamore Pool & Spa", "sycamorepoolspa.example", "(925) 555-0163", "", "", "Danville, CA 94526", "", "", "IDEAL", "N", "Not contacted"),
        ("Landscaping", "Westside Garden Design", "westsidegardendesign.example", "(925) 555-0171", "", "", "Danville, CA 94526", "", "", "STRONG", "N", "Not contacted"),
        ("General Dentistry", "Hartz Avenue Dental Group", "hartzdental.example", "(925) 555-0181", "", "", "Danville, CA 94526", "", "", "IDEAL", "N", "Not contacted"),
        ("Pediatric Dentistry", "Diablo Pediatric Smiles", "diablopedsmiles.example", "(925) 555-0184", "", "", "Danville, CA 94526", "", "", "STRONG", "N", "Not contacted"),
        ("Plumbing", "Round Hill Plumbing Co.", "roundhillplumbing.example", "(925) 555-0193", "", "", "Danville, CA 94526", "", "", "GOOD", "N", "Not contacted"),
        ("Roofing", "San Ramon Valley Roofing", "srvroofing.example", "(925) 555-0201", "", "", "Danville, CA 94526", "", "", "STRONG", "N", "Not contacted"),
        ("House Cleaning", "Pure Home Cleaners (Tri-Valley)", "puretrivalley.example", "(925) 555-0211", "", "", "Danville, CA 94526", "", "", "GOOD", "N", "Not contacted"),
        ("Med Spa / Aesthetics", "Diablo Aesthetics & Wellness", "diabloaesthetics.example", "(925) 555-0221", "", "", "Danville, CA 94526", "", "", "IDEAL", "N", "Not contacted"),
    ],
    "San Ramon": [
        ("HVAC", "Bishop Ranch HVAC Services", "brhvac.example", "(925) 555-0240", "", "", "San Ramon, CA 94583", "", "", "IDEAL", "N", "Not contacted"),
        ("HVAC", "Dougherty Valley Heating", "doughertyhvac.example", "(925) 555-0241", "", "", "San Ramon, CA 94582", "", "", "STRONG", "N", "Not contacted"),
        ("Pest Control", "Tri-Valley Pest Pros", "tvpestpros.example", "(925) 555-0252", "", "", "San Ramon, CA 94583", "", "", "STRONG", "N", "Not contacted"),
        ("Pool Service", "Gale Ranch Pool Care", "galeranchpool.example", "(925) 555-0263", "", "", "San Ramon, CA 94582", "", "", "IDEAL", "N", "Not contacted"),
        ("Landscaping", "Windemere Landscape Co.", "windemerelandscape.example", "(925) 555-0271", "", "", "San Ramon, CA 94582", "", "", "STRONG", "N", "Not contacted"),
        ("General Dentistry", "Crow Canyon Family Dentistry", "ccfamilydental.example", "(925) 555-0281", "", "", "San Ramon, CA 94583", "", "", "IDEAL", "N", "Not contacted"),
        ("Pediatric Dentistry", "Little Ranch Pediatric Dentistry", "littleranchpeds.example", "(925) 555-0284", "", "", "San Ramon, CA 94582", "", "", "STRONG", "N", "Not contacted"),
        ("Plumbing", "Twin Creeks Plumbing", "twincreeksplumbing.example", "(925) 555-0291", "", "", "San Ramon, CA 94583", "", "", "GOOD", "N", "Not contacted"),
        ("House Cleaning", "Dougherty Home Cleaners", "doughertyhomeclean.example", "(925) 555-0301", "", "", "San Ramon, CA 94582", "", "", "GOOD", "N", "Not contacted"),
        ("Med Spa / Aesthetics", "San Ramon Med Spa & Laser", "srmedspa.example", "(925) 555-0311", "", "", "San Ramon, CA 94583", "", "", "IDEAL", "N", "Not contacted"),
        ("Roofing", "Ridgeline Roofing of San Ramon", "ridgelineroofing.example", "(925) 555-0321", "", "", "San Ramon, CA 94583", "", "", "STRONG", "N", "Not contacted"),
    ],
}


def style_header_row(ws, row, columns):
    for col_idx, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[row].height = 28


def add_prospects_sheet(wb, community_name, zip_code, seed):
    ws = wb.create_sheet(title=f"{community_name} {zip_code.split('/')[0]}")

    title = ws.cell(row=1, column=1,
                    value=f"{community_name}, CA — {zip_code} · Prospect Tracker")
    title.font = Font(bold=True, size=16, color="0E1F3A")
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(PROSPECT_COLUMNS))
    ws.row_dimensions[1].height = 26

    sub = ws.cell(row=2, column=1,
                  value="One business per category gets the slot. Mark Locked? = Y once deposit is paid.")
    sub.font = Font(italic=True, color="6A7080", size=10)
    ws.merge_cells(start_row=2, start_column=1,
                   end_row=2, end_column=len(PROSPECT_COLUMNS))

    style_header_row(ws, 4, PROSPECT_COLUMNS)

    row = 5
    for prospect in seed:
        for col_idx, value in enumerate(prospect, start=1):
            c = ws.cell(row=row, column=col_idx, value=value)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
        row += 1

    blanks_to_add = max(0, 30 - len(seed))
    for _ in range(blanks_to_add):
        for col_idx, _ in enumerate(PROSPECT_COLUMNS, start=1):
            c = ws.cell(row=row, column=col_idx, value=None)
            c.border = BORDER
        row += 1

    ws.freeze_panes = "A5"

    rating_dv = DataValidation(type="list",
                               formula1='"' + ",".join(RATINGS) + '"',
                               allow_blank=True)
    status_dv = DataValidation(type="list",
                               formula1='"' + ",".join(STATUS) + '"',
                               allow_blank=True)
    locked_dv = DataValidation(type="list",
                               formula1='"Y,N"',
                               allow_blank=True)
    deposit_dv = DataValidation(type="list",
                                formula1='"Y,N"',
                                allow_blank=True)
    cat_dv = DataValidation(type="list",
                            formula1='"' + ",".join(CATEGORIES) + '"',
                            allow_blank=True)

    ws.add_data_validation(rating_dv)
    ws.add_data_validation(status_dv)
    ws.add_data_validation(locked_dv)
    ws.add_data_validation(deposit_dv)
    ws.add_data_validation(cat_dv)

    last = row - 1
    cat_dv.add(f"A5:A{last}")
    rating_dv.add(f"J5:J{last}")
    locked_dv.add(f"K5:K{last}")
    status_dv.add(f"L5:L{last}")
    deposit_dv.add(f"P5:P{last}")


def add_summary_sheet(wb, community_name, zip_code):
    ws = wb.create_sheet(title="Campaign Summary", index=0)

    title = ws.cell(row=1, column=1,
                    value=f"{community_name} {zip_code} · Campaign Summary")
    title.font = Font(bold=True, size=18, color="0E1F3A")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

    rows = [
        ("Community", community_name),
        ("Zip code(s)", zip_code),
        ("Target households", 6000),
        ("Target ad slots", 8),
        ("Standard ad price ($)", 650),
        ("Premium ad price ($)", 950),
        ("Printing cost est. ($)", 350),
        ("EDDM postage rate ($/piece)", 0.247),
        ("Misc cost est. ($)", 100),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        kc = ws.cell(row=i, column=1, value=k)
        vc = ws.cell(row=i, column=2, value=v)
        kc.font = Font(bold=True, color="0E1F3A")
        vc.alignment = Alignment(horizontal="left")

    ws.cell(row=13, column=1, value="EDDM postage cost ($)").font = Font(bold=True)
    ws.cell(row=13, column=2, value="=B5*B10")
    ws.cell(row=14, column=1, value="Total cost est. ($)").font = Font(bold=True)
    ws.cell(row=14, column=2, value="=B9+B13+B11")
    ws.cell(row=15, column=1, value="Revenue if 6 standard ($)").font = Font(bold=True)
    ws.cell(row=15, column=2, value="=B7*6")
    ws.cell(row=16, column=1, value="Revenue if 6 std + 2 premium ($)").font = Font(bold=True)
    ws.cell(row=16, column=2, value="=B7*6+B8*2")
    ws.cell(row=17, column=1, value="Net @ 6 standard ($)").font = Font(bold=True)
    ws.cell(row=17, column=2, value="=B15-B14")
    ws.cell(row=18, column=1, value="Net @ 6 + 2 premium ($)").font = Font(bold=True)
    ws.cell(row=18, column=2, value="=B16-B14")

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16


def build_workbook(community_name, zip_code, output_filename):
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    add_summary_sheet(wb, community_name, zip_code)
    add_prospects_sheet(wb, community_name, zip_code,
                        SEED_PROSPECTS.get(community_name, []))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / output_filename
    wb.save(out)
    print(f"Wrote {out}")


def main():
    build_workbook("Danville", "94526", "prospect-tracker_danville-94526.xlsx")
    build_workbook("San Ramon", "94582/94583", "prospect-tracker_san-ramon-94582-94583.xlsx")


if __name__ == "__main__":
    main()
